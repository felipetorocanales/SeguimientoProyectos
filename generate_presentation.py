import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_presentation(output_path):
    prs = Presentation()
    # 16:9 widescreen slides
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank_layout = prs.slide_layouts[6]

    # Color Palette: Nexus Modern Executive Dark Theme
    COLOR_BG = RGBColor(6, 19, 48)          # #061330
    COLOR_CARD = RGBColor(12, 31, 75)       # #0c1f4b
    COLOR_CARD_BORDER = RGBColor(30, 64, 130)
    COLOR_PRIMARY = RGBColor(56, 189, 248)  # #38bdf8 (Sky Blue)
    COLOR_ACCENT = RGBColor(99, 102, 241)   # #6366f1 (Indigo/Purple)
    COLOR_EMERALD = RGBColor(16, 185, 129)  # #10b981 (Success/Resume)
    COLOR_AMBER = RGBColor(245, 158, 11)    # #f59e0b (Pause/Warning)
    COLOR_CORAL = RGBColor(248, 113, 113)   # #f87171 (Soft Red)
    COLOR_TELEGRAM = RGBColor(41, 169, 234) # #29a9ea (Telegram Blue)
    COLOR_TEXT_WHITE = RGBColor(248, 250, 252)
    COLOR_TEXT_MUTED = RGBColor(148, 163, 184) # #94a3b8

    def set_slide_background(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = COLOR_BG
        bg.line.fill.background()

        top_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.12))
        top_bar.fill.solid()
        top_bar.fill.fore_color.rgb = COLOR_PRIMARY
        top_bar.line.fill.background()

    def add_header(slide, tag_text, title_text, subtitle_text=""):
        badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(0.55), Inches(3.2), Inches(0.35))
        badge.fill.solid()
        badge.fill.fore_color.rgb = COLOR_CARD
        badge.line.color.rgb = COLOR_PRIMARY
        badge.line.width = Pt(1)
        tf_b = badge.text_frame
        tf_b.word_wrap = False
        p_b = tf_b.paragraphs[0]
        p_b.text = tag_text.upper()
        p_b.font.size = Pt(10)
        p_b.font.bold = True
        p_b.font.color.rgb = COLOR_PRIMARY
        p_b.alignment = PP_ALIGN.CENTER

        tx_box = slide.shapes.add_textbox(Inches(1.0), Inches(0.92), Inches(11.3), Inches(1.0))
        tf = tx_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.size = Pt(25)
        p.font.bold = True
        p.font.color.rgb = COLOR_TEXT_WHITE

        if subtitle_text:
            p2 = tf.add_paragraph()
            p2.text = subtitle_text
            p2.font.size = Pt(13)
            p2.font.color.rgb = COLOR_TEXT_MUTED
            p2.space_before = Pt(4)

    # ═════════════════════════════════════════════════════════════════════════
    # SLIDE 1: Portada
    # ═════════════════════════════════════════════════════════════════════════
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1)

    center_card = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), Inches(1.3), Inches(10.333), Inches(4.9))
    center_card.fill.solid()
    center_card.fill.fore_color.rgb = COLOR_CARD
    center_card.line.color.rgb = COLOR_PRIMARY
    center_card.line.width = Pt(1.5)

    tf1 = center_card.text_frame
    tf1.word_wrap = True
    
    p = tf1.paragraphs[0]
    p.text = "NEXUS TRACKER"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY
    p.alignment = PP_ALIGN.CENTER
    p.space_before = Pt(25)

    p2 = tf1.add_paragraph()
    p2.text = "Capacitación Nuevas Funcionalidades"
    p2.font.size = Pt(32)
    p2.font.bold = True
    p2.font.color.rgb = COLOR_TEXT_WHITE
    p2.alignment = PP_ALIGN.CENTER
    p2.space_before = Pt(10)

    p3 = tf1.add_paragraph()
    p3.text = "Ciclo Único Variable · Gestión de Pausas · Cronograma Flexible · Bot de Telegram"
    p3.font.size = Pt(14)
    p3.font.color.rgb = COLOR_TEXT_MUTED
    p3.alignment = PP_ALIGN.CENTER
    p3.space_before = Pt(12)

    p4 = tf1.add_paragraph()
    p4.text = "Experis Chile · Sistema de Seguimiento Estratégico · Versión Producción 2026"
    p4.font.size = Pt(12)
    p4.font.color.rgb = RGBColor(96, 165, 250)
    p4.alignment = PP_ALIGN.CENTER
    p4.space_before = Pt(35)

    # ═════════════════════════════════════════════════════════════════════════
    # SLIDE 2: Agenda de Capacitación
    # ═════════════════════════════════════════════════════════════════════════
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2)
    add_header(s2, "Agenda de Capacitación", "¿Qué veremos hoy?", "Evolución integral del modelo de seguimiento hacia mayor dinamismo y flexibilidad.")

    cards_data_s2 = [
        ("1. Ciclo Único Variable", "Adiós a fases rígidas estandarizadas: cada proyecto es un ciclo fluido gobernado por su bitácora de comentarios.", COLOR_ACCENT),
        ("2. Pausas & Extensión Dinámica", "Pausar/reanudar con fecha flexible y justificación obligatoria, sumando días sin perder la meta inicial.", COLOR_AMBER),
        ("3. Cronograma Semanas / Meses", "Suavizado visual de aplazamientos (sin alarmismo) y selector [Semanas | Meses] individual por proyecto.", COLOR_PRIMARY),
        ("4. Bot de Telegram Oficial", "Comandos de visualización ejecutiva (/resumen, /proyectos, /buscar) y actualización guiada (/actualizar).", COLOR_TELEGRAM)
    ]

    for i, (ctitle, cdesc, ccolor) in enumerate(cards_data_s2):
        col = i % 2
        row = i // 2
        left = Inches(1.0 + col * 5.8)
        top = Inches(2.2 + row * 2.3)
        
        card = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, Inches(5.5), Inches(2.0))
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_CARD
        card.line.color.rgb = ccolor
        card.line.width = Pt(1.5)

        tf = card.text_frame
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.text = ctitle
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = ccolor
        
        p2 = tf.add_paragraph()
        p2.text = cdesc
        p2.font.size = Pt(13)
        p2.font.color.rgb = COLOR_TEXT_WHITE
        p2.space_before = Pt(8)

    # ═════════════════════════════════════════════════════════════════════════
    # SLIDE 3: Transición al Ciclo Único Variable
    # ═════════════════════════════════════════════════════════════════════════
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3)
    add_header(s3, "Evolución del Modelo", "Transición: De Fases Estandarizadas a Ciclo Único Variable", "Simplificación operativa donde la realidad del proyecto guía el flujo de trabajo.")

    left_old = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(2.1), Inches(5.5), Inches(4.7))
    left_old.fill.solid()
    left_old.fill.fore_color.rgb = COLOR_CARD
    left_old.line.color.rgb = COLOR_CORAL
    left_old.line.width = Pt(1.5)
    tf_o = left_old.text_frame
    tf_o.word_wrap = True

    p = tf_o.paragraphs[0]
    p.text = "❌ Modelo Anterior: Fases Rígidas"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = COLOR_CORAL

    bullets_old = [
        "Estructura predefinida obligatoria (Levantamiento, Diseño, Desarrollo, Pruebas, Despliegue).",
        "Burocracia de actualización: Había que editar fechas y porcentajes fase por fase individualmente.",
        "Desalineación con la realidad: Muchos proyectos no siguen cascadas tradicionales o tienen ciclos ágiles híbridos.",
        "Confusión en fechas: Fechas intermedias ficticias que no representaban el avance real del cliente."
    ]
    for b in bullets_old:
        p_b = tf_o.add_paragraph()
        p_b.text = f"•  {b}"
        p_b.font.size = Pt(12)
        p_b.font.color.rgb = COLOR_TEXT_MUTED
        p_b.space_before = Pt(8)

    right_new = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(2.1), Inches(5.5), Inches(4.7))
    right_new.fill.solid()
    right_new.fill.fore_color.rgb = COLOR_CARD
    right_new.line.color.rgb = COLOR_EMERALD
    right_new.line.width = Pt(1.5)
    tf_n = right_new.text_frame
    tf_n.word_wrap = True

    p = tf_n.paragraphs[0]
    p.text = "✅ Nuevo Modelo: Ciclo Único Variable"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = COLOR_EMERALD

    bullets_new = [
        "Un solo horizonte claro por proyecto: Fecha de Inicio → Fecha de Compromiso de Entrega.",
        "Gobierno por Comentarios / Bitácora: El avance y estado se van construyendo dinámicamente con las actualizaciones que ingresa el equipo.",
        "Motor de Inteligencia Artificial: Infiere la etapa táctica real a partir del contexto de los comentarios ingresados.",
        "Agilidad total: Actualizar el avance real (% editable directo) y reportar hitos en segundos desde la app o Telegram."
    ]
    for b in bullets_new:
        p_b = tf_n.add_paragraph()
        p_b.text = f"•  {b}"
        p_b.font.size = Pt(12)
        p_b.font.color.rgb = COLOR_TEXT_WHITE
        p_b.space_before = Pt(8)

    # ═════════════════════════════════════════════════════════════════════════
    # SLIDE 4: Dinamismo Basado en Comentarios e IA
    # ═════════════════════════════════════════════════════════════════════════
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_background(s4)
    add_header(s4, "Gestión Viva del Proyecto", "¿Cómo Opera el Ciclo Variable Basado en Comentarios?", "Cada comentario ingresado enriquece el estado del proyecto y ajusta la visión en tiempo real.")

    cycle_steps = [
        ("1. Reporte Rápido y Flexible", "El líder o desarrollador escribe notas directas: avances, acuerdos con el cliente, bloqueos o hitos completados.", COLOR_PRIMARY),
        ("2. Inferencia Inteligente (IA)", "El sistema analiza semánticamente las notas para deducir la etapa activa (ej: 'Pruebas UAT', 'Despliegue', 'Afinamiento').", COLOR_ACCENT),
        ("3. Indicadores de Salud", "Detecta alertas tempranas (🟢 A Tiempo, 🟡 En Riesgo, 🚨 Retrasado) según el gap entre tiempo hábil y avance reportado.", COLOR_AMBER)
    ]

    for i, (stitle, sdesc, scolor) in enumerate(cycle_steps):
        card = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0 + i * 3.9), Inches(2.2), Inches(3.6), Inches(4.6))
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_CARD
        card.line.color.rgb = scolor
        card.line.width = Pt(1.5)

        tf = card.text_frame
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.text = stitle
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = scolor

        p2 = tf.add_paragraph()
        p2.text = sdesc
        p2.font.size = Pt(13)
        p2.font.color.rgb = COLOR_TEXT_WHITE
        p2.space_before = Pt(14)

    # ═════════════════════════════════════════════════════════════════════════
    # SLIDE 5: Sistema de Pausas de Proyectos
    # ═════════════════════════════════════════════════════════════════════════
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_background(s5)
    add_header(s5, "Módulo de Pausas", "Pausar y Reanudar Proyectos de Forma Controlada", "Detiene el conteo del tiempo hábil sin alterar las métricas de avance de los desarrolladores.")

    left_card = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(2.2), Inches(5.5), Inches(4.6))
    left_card.fill.solid()
    left_card.fill.fore_color.rgb = COLOR_CARD
    left_card.line.color.rgb = COLOR_CARD_BORDER
    tf = left_card.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = "⏸️ ¿Cómo Funciona la Pausa?"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = COLOR_AMBER

    bullets_left = [
        "Botón 'Pausar' disponible en la tarjeta de cada proyecto para Líderes y Administradores.",
        "Fecha de inicio de pausa flexible: Propone hoy por defecto pero permite indicar una fecha anterior si la detención ocurrió antes.",
        "Motivo Obligatorio: Es mandatorio ingresar la justificación (ej: 'Espera de credenciales', 'Validación cliente', 'Bloqueo presupuestario').",
        "Estado Visual Inmediato: Píldora distintiva '⏸️ En Pausa (+Xd)' y filtrado directo en el tablero."
    ]
    for b in bullets_left:
        p_b = tf.add_paragraph()
        p_b.text = f"•  {b}"
        p_b.font.size = Pt(12)
        p_b.font.color.rgb = COLOR_TEXT_WHITE
        p_b.space_before = Pt(8)

    right_card = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(2.2), Inches(5.5), Inches(4.6))
    right_card.fill.solid()
    right_card.fill.fore_color.rgb = COLOR_CARD
    right_card.line.color.rgb = COLOR_CARD_BORDER
    tf_r = right_card.text_frame
    tf_r.word_wrap = True

    p = tf_r.paragraphs[0]
    p.text = "▶️ Reanudación y Cierre de Pausa"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = COLOR_EMERALD

    bullets_right = [
        "El botón conmuta a '▶️ Reanudar': Un solo clic reincorpora el proyecto al flujo activo de trabajo.",
        "Cálculo exacto de días: Computa con precisión los días que estuvo pausado y cierra el intervalo formalmente.",
        "Historial Completo: Cada proyecto almacena todas sus pausas con fechas exactas, usuario responsable y motivos.",
        "Pausas Múltiples: Un proyecto puede pausarse y reactivarse reiteradas veces; todos los días se totalizan automáticamente."
    ]
    for b in bullets_right:
        p_b = tf_r.add_paragraph()
        p_b.text = f"•  {b}"
        p_b.font.size = Pt(12)
        p_b.font.color.rgb = COLOR_TEXT_WHITE
        p_b.space_before = Pt(8)

    # ═════════════════════════════════════════════════════════════════════════
    # SLIDE 6: Extensión Dinámica y Fechas
    # ═════════════════════════════════════════════════════════════════════════
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_background(s6)
    add_header(s6, "Lógica de Fechas", "Extensión Dinámica de la Fecha de Entrega", "Garantía de que las pausas por terceros no impacten negativamente la evaluación del equipo.")

    steps_ext = [
        ("1. Meta Inicial Inmutable", "La fecha de compromiso original (ej: 03/02/2027) jamás se borra. Queda protegida como 'Meta Inicial' con candado de auditoría.", COLOR_PRIMARY),
        ("2. Suma Día tras Día", "Mientras el proyecto permanezca en pausa, el sistema suma los días a la fecha de entrega final, aplazándola dinámicamente en tiempo real.", COLOR_AMBER),
        ("3. Desfase Justificado", "Al reanudar, la nueva fecha de entrega queda formalmente consolidada sumando los días efectivos de retraso justificado.", COLOR_EMERALD)
    ]

    for i, (stitle, sdesc, scolor) in enumerate(steps_ext):
        card = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0 + i * 3.9), Inches(2.2), Inches(3.6), Inches(4.6))
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_CARD
        card.line.color.rgb = scolor
        card.line.width = Pt(1.5)

        tf = card.text_frame
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.text = stitle
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = scolor

        p2 = tf.add_paragraph()
        p2.text = sdesc
        p2.font.size = Pt(13)
        p2.font.color.rgb = COLOR_TEXT_WHITE
        p2.space_before = Pt(14)

    # ═════════════════════════════════════════════════════════════════════════
    # SLIDE 7: Mejoras Visuales en el Cronograma (Gantt)
    # ═════════════════════════════════════════════════════════════════════════
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_background(s7)
    add_header(s7, "Experiencia Visual", "Diseño Suave y Amigable del Cronograma", "Alineado con el feedback del equipo: eliminar sensaciones de alarma innecesarias.")

    card_gantt_1 = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(2.2), Inches(5.5), Inches(4.6))
    card_gantt_1.fill.solid()
    card_gantt_1.fill.fore_color.rgb = COLOR_CARD
    card_gantt_1.line.color.rgb = COLOR_CARD_BORDER
    tf = card_gantt_1.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "🎨 Tono Suave de Días Aplazados"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY

    b_list1 = [
        "Antes: Franjas naranjas/amarillas brillantes de advertencia que causaban alerta visual excesiva.",
        "Ahora: Estilo translúcido elegante en cian y azul noche con borde punteado suave.",
        "Indicador discreto: Píldora compacta '⏳ +Xd (DD/MM/AAAA)' integrada con armonía en la barra del proyecto.",
        "Línea de Compromiso Inicial: Delgada marca luminosa que señala con claridad dónde terminaba la meta original."
    ]
    for b in b_list1:
        p_b = tf.add_paragraph()
        p_b.text = f"•  {b}"
        p_b.font.size = Pt(12)
        p_b.font.color.rgb = COLOR_TEXT_WHITE
        p_b.space_before = Pt(8)

    card_gantt_2 = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(2.2), Inches(5.5), Inches(4.6))
    card_gantt_2.fill.solid()
    card_gantt_2.fill.fore_color.rgb = COLOR_CARD
    card_gantt_2.line.color.rgb = COLOR_CARD_BORDER
    tf2 = card_gantt_2.text_frame
    tf2.word_wrap = True

    p = tf2.paragraphs[0]
    p.text = "⏸️ Visualización de Pausas en Barra"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = COLOR_AMBER

    b_list2 = [
        "Espacio de Pausa: Se proyecta una franja translúcida rayada sobre el periodo exacto en que ocurrió la detención.",
        "Ícono de Pausa: Marcador '⏸️' con tooltip detallado que muestra fechas exactas y motivo.",
        "Sin alteraciones de avance: El avance real completado (% real) se mantiene intacto.",
        "Transparencia para clientes y gerencia: Evidencia rápida de por qué un hito se movió en el tiempo."
    ]
    for b in b_list2:
        p_b = tf2.add_paragraph()
        p_b.text = f"•  {b}"
        p_b.font.size = Pt(12)
        p_b.font.color.rgb = COLOR_TEXT_WHITE
        p_b.space_before = Pt(8)

    # ═════════════════════════════════════════════════════════════════════════
    # SLIDE 8: Visión Semanas / Meses en Cada Proyecto
    # ═════════════════════════════════════════════════════════════════════════
    s8 = prs.slides.add_slide(blank_layout)
    set_slide_background(s8)
    add_header(s8, "Navegación Temporal", "Visión por Semanas y por Meses en Cada Proyecto", "Flexibilidad para analizar proyectos de corto o largo plazo con la escala adecuada.")

    card_feat = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(2.2), Inches(11.333), Inches(4.6))
    card_feat.fill.solid()
    card_feat.fill.fore_color.rgb = COLOR_CARD
    card_feat.line.color.rgb = COLOR_PRIMARY
    card_feat.line.width = Pt(1.5)
    
    tf = card_feat.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "📅 Control de Escala Temporal Individual"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY

    bullets_s8 = [
        "Selector [ Semanas | Meses ] en la cabecera del cronograma de cada proyecto: Ya no dependes únicamente de la escala global.",
        "Vista Mensual para proyectos extensos: Permite visualizar planificaciones semestrales o anuales completas sin scroll horizontal excesivo.",
        "Vista Semanal para seguimiento táctico: Permite revisar entregables inmediatos, fechas de corte y sprints con máxima granularidad.",
        "Independencia entre proyectos: Puedes tener un proyecto en vista mensual y otro en vista semanal al mismo tiempo sin que interfieran.",
        "Línea de 'Hoy' adaptativa: La guía vertical roja se posiciona automáticamente en la semana o en la proporción exacta del mes actual."
    ]
    for b in bullets_s8:
        p_b = tf.add_paragraph()
        p_b.text = f"✔   {b}"
        p_b.font.size = Pt(13)
        p_b.font.color.rgb = COLOR_TEXT_WHITE
        p_b.space_before = Pt(10)

    # ═════════════════════════════════════════════════════════════════════════
    # SLIDE 9: Bot de Telegram - Comandos de Visualización y Consulta (NUEVA)
    # ═════════════════════════════════════════════════════════════════════════
    s9 = prs.slides.add_slide(blank_layout)
    set_slide_background(s9)
    add_header(s9, "Integración Telegram", "Bot de Telegram: Comandos de Consulta y Estado", "Seguimiento ágil en tiempo real desde dispositivos móviles sin entrar al navegador.")

    tele_commands = [
        ("📊 /resumen", "KPIs Ejecutivos en Segundos", "Total proyectos, completados, a tiempo 🟢, en riesgo ⚠️, retrasados 🚨, % SLA y lista de proyectos críticos en foco.", COLOR_PRIMARY),
        ("📋 /proyectos", "Listado con Filtros Inteligentes", "Lista proyectos activos con tiempo transcurrido y salud.\nAdmite filtros: /proyectos todo, completados, riesgo, retrasados.", COLOR_ACCENT),
        ("🔍 /buscar <nombre>", "Ficha Completa del Proyecto", "Detalle de cliente, responsable, fecha de entrega, salud actual y las últimas 4 actualizaciones de la bitácora.", COLOR_AMBER)
    ]

    for i, (tcmd, ttitle, tdesc, tcol) in enumerate(tele_commands):
        card = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0 + i * 3.9), Inches(2.2), Inches(3.6), Inches(4.6))
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_CARD
        card.line.color.rgb = tcol
        card.line.width = Pt(1.5)

        tf = card.text_frame
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.text = tcmd
        p.font.size = Pt(17)
        p.font.bold = True
        p.font.color.rgb = tcol

        p2 = tf.add_paragraph()
        p2.text = ttitle
        p2.font.size = Pt(13)
        p2.font.bold = True
        p2.font.color.rgb = COLOR_TEXT_WHITE
        p2.space_before = Pt(8)

        p3 = tf.add_paragraph()
        p3.text = tdesc
        p3.font.size = Pt(11)
        p3.font.color.rgb = COLOR_TEXT_MUTED
        p3.space_before = Pt(8)

    # ═════════════════════════════════════════════════════════════════════════
    # SLIDE 10: Bot de Telegram - Flujo de Actualización /actualizar (NUEVA)
    # ═════════════════════════════════════════════════════════════════════════
    s10 = prs.slides.add_slide(blank_layout)
    set_slide_background(s10)
    add_header(s10, "Integración Telegram", "Bot de Telegram: Asistente Interactivo /actualizar", "Reporte asistido paso a paso con botones dinámicos (Inline Keyboards).")

    wizard_steps = [
        ("Paso 1: /actualizar", "Selección de Proyecto", "Despliega una lista de botones interactivos con todos los proyectos activos y su indicador de salud (🟢, ⚠️, 🚨).", COLOR_PRIMARY),
        ("Paso 2: Logros", "Reporte de Avances", "Solicita qué hitos o tareas se completaron en el período. Opción rápida: 'Omitir logros'.", COLOR_EMERALD),
        ("Paso 3: Bloqueos", "Frenos o Impedimentos", "Pregunta qué los está deteniendo (accesos, feedback cliente, bugs). Opción rápida: 'Ninguno (Todo fluye)'.", COLOR_AMBER),
        ("Paso 4: IA & Bitácora", "Impacto Inmediato", "Guarda la entrada fechada en Firestore (DD/MM/AA: Logros | Bloqueos), actualiza la web y dispara la inferencia de IA.", COLOR_ACCENT)
    ]

    for i, (wstep, wtitle, wdesc, wcol) in enumerate(wizard_steps):
        left_pos = Inches(1.0 + i * 2.88)
        card = s10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_pos, Inches(2.2), Inches(2.7), Inches(4.6))
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_CARD
        card.line.color.rgb = wcol
        card.line.width = Pt(1.5)

        tf = card.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = wstep
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = wcol

        p2 = tf.add_paragraph()
        p2.text = wtitle
        p2.font.size = Pt(12)
        p2.font.bold = True
        p2.font.color.rgb = COLOR_TEXT_WHITE
        p2.space_before = Pt(6)

        p3 = tf.add_paragraph()
        p3.text = wdesc
        p3.font.size = Pt(11)
        p3.font.color.rgb = COLOR_TEXT_MUTED
        p3.space_before = Pt(8)

    # ═════════════════════════════════════════════════════════════════════════
    # SLIDE 11: Buenas Prácticas y Trazabilidad
    # ═════════════════════════════════════════════════════════════════════════
    s11 = prs.slides.add_slide(blank_layout)
    set_slide_background(s11)
    add_header(s11, "Buenas Prácticas", "Trazabilidad, Bitácora y Recomendaciones de Uso", "Lineamientos para que los líderes y gerentes mantengan la data confiable y transparente.")

    tips_data = [
        ("Reportar por Telegram o Web", "Usa /actualizar al salir de reuniones o terminar sprints. Entre más frecuente el reporte, más certera es la IA en la inferencia.", COLOR_ACCENT),
        ("Pausar a Tiempo", "Si el proyecto se detiene por falta de respuesta del cliente o bloqueos externos, registrar la pausa de inmediato para no penalizar el tiempo en Dashboard.", COLOR_AMBER),
        ("Fecha Retroactiva Responsable", "Si por alguna razón no pudiste pausar el mismo día, utiliza el selector de fecha de inicio de pausa para ingresar la fecha real del bloqueo.", COLOR_PRIMARY),
        ("Auditoría y Comités", "Toda la bitácora alimenta de forma automática los logs de auditoría y el resumen ejecutivo de IA para los comités semanales de gerencia.", COLOR_EMERALD)
    ]

    for i, (ttitle, tdesc, tcolor) in enumerate(tips_data):
        col = i % 2
        row = i // 2
        left = Inches(1.0 + col * 5.8)
        top = Inches(2.2 + row * 2.3)
        
        card = s11.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, Inches(5.5), Inches(2.0))
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_CARD
        card.line.color.rgb = tcolor
        card.line.width = Pt(1.5)

        tf = card.text_frame
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.text = ttitle
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = tcolor
        
        p2 = tf.add_paragraph()
        p2.text = tdesc
        p2.font.size = Pt(12)
        p2.font.color.rgb = COLOR_TEXT_WHITE
        p2.space_before = Pt(6)

    # ═════════════════════════════════════════════════════════════════════════
    # SLIDE 12: Cierre & Preguntas
    # ═════════════════════════════════════════════════════════════════════════
    s12 = prs.slides.add_slide(blank_layout)
    set_slide_background(s12)

    card_end = s12.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), Inches(1.4), Inches(10.333), Inches(4.8))
    card_end.fill.solid()
    card_end.fill.fore_color.rgb = COLOR_CARD
    card_end.line.color.rgb = COLOR_PRIMARY
    card_end.line.width = Pt(1.5)

    tf_e = card_end.text_frame
    tf_e.word_wrap = True
    
    p = tf_e.paragraphs[0]
    p.text = "¿PREGUNTAS O COMENTARIOS?"
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY
    p.alignment = PP_ALIGN.CENTER
    p.space_before = Pt(35)

    p2 = tf_e.add_paragraph()
    p2.text = "Las nuevas funcionalidades ya se encuentran disponibles y activas en producción."
    p2.font.size = Pt(16)
    p2.font.color.rgb = COLOR_TEXT_WHITE
    p2.alignment = PP_ALIGN.CENTER
    p2.space_before = Pt(14)

    p3 = tf_e.add_paragraph()
    p3.text = "🌐 Plataforma Web: https://nexus-tracker-b7a75.web.app\n🤖 Bot de Telegram: /start · /resumen · /actualizar"
    p3.font.size = Pt(15)
    p3.font.bold = True
    p3.font.color.rgb = COLOR_EMERALD
    p3.alignment = PP_ALIGN.CENTER
    p3.space_before = Pt(18)

    p4 = tf_e.add_paragraph()
    p4.text = "Gracias por su compromiso con la excelencia operativa · Experis Chile"
    p4.font.size = Pt(12)
    p4.font.color.rgb = COLOR_TEXT_MUTED
    p4.alignment = PP_ALIGN.CENTER
    p4.space_before = Pt(22)

    prs.save(output_path)
    print(f"Presentation saved successfully to {output_path}")

if __name__ == "__main__":
    out_dir = r"c:\Users\Felipe\Desktop\SeguimientoProyectos"
    out_file = os.path.join(out_dir, "Capacitacion_Nexus_Tracker_Ciclo_Unico.pptx")
    try:
        create_presentation(out_file)
    except PermissionError:
        out_file_alt = os.path.join(out_dir, "Capacitacion_Nexus_Tracker_v3.pptx")
        create_presentation(out_file_alt)
